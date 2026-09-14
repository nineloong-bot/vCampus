# -*- coding: utf-8 -*-
"""将《章节/*.md》组装为 Word 软件设计说明书。

用法：
    python 工具/build_doc.py

输入：设计说明书素材/章节/*.md （按文件名排序），以及 设计说明书素材/图/*.png
输出：D:\\_store\\vCampus\\软件设计说明书-vCampus.docx

支持的 Markdown 子集：
    # / ## / ### / ####  标题
    | a | b |            表格（第二行为 --- 分隔行）
    - 项目 / 1. 项目      列表
    > 说明                引用块（楷体缩进）
    ![说明](图/xxx.png)   居中插图 + 图题
    普通段落              正文
    行内 **加粗**、`等宽`
"""
from __future__ import annotations

import glob
import os
import re
import sys
import unicodedata

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))       # 设计说明书素材/
ROOT = os.path.dirname(os.path.dirname(BASE))                            # 仓库根
CHAPTER_DIR = os.path.join(BASE, "章节")
IMG_DIR = os.path.join(BASE, "图")
DEMO = (r"C:\Users\31966\.dsh\attachments\v1\files\14"
        r"\147a3454e4bb2a72e4bdbbfc749c2d78116c7267e05fdb90086dc917bfc70ae5"
        r"\软件设计说明书DEMO(20250825).docx")
OUT = os.path.join(ROOT, "软件设计说明书-vCampus.docx")
if len(sys.argv) > 1:
    OUT = os.path.abspath(sys.argv[1])

FONT_ASCII = "Times New Roman"
FONT_EA = "宋体"
FONT_EA_HEAD = "黑体"
CONTENT_WIDTH_CM = 16.0
PAGE_MARGIN_CM = 2.4

CODE_RE = re.compile(r"`([^`]+)`")
BOLD_RE = re.compile(r"\*\*([^*]+)\*\*")
IMG_RE = re.compile(r"^!\[(.*?)\]\((.+?)\)\s*$")
TABLE_SEP_RE = re.compile(r"^\|[\s:\-|]+\|$")


# ----------------------------------------------------------------- 低层工具
def _el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k), v)
    return e


def set_style_font(style, ascii_font=FONT_ASCII, ea_font=FONT_EA,
                   size=None, bold=None, color=None):
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = _el("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), ascii_font)
    rfonts.set(qn("w:hAnsi"), ascii_font)
    rfonts.set(qn("w:eastAsia"), ea_font)
    if size is not None:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if color is not None:
        style.font.color.rgb = RGBColor(*color)


def ensure_heading3(doc):
    """DEMO 模板只定义了 Heading 1/2，这里补出 Heading 3。"""
    try:
        st = doc.styles["Heading 3"]
        if st.name:
            return st
    except KeyError:
        pass
    from docx.enum.style import WD_STYLE_TYPE
    st = doc.styles.add_style("Heading 3", WD_STYLE_TYPE.PARAGRAPH)
    st.base_style = doc.styles["Heading 2"]
    return st


def add_field(paragraph, instr: str):
    r = paragraph.add_run()
    r._r.append(_el("w:fldChar", **{"w:fldCharType": "begin"}))
    r2 = paragraph.add_run()
    it = _el("w:instrText", **{"xml:space": "preserve"})
    it.text = instr
    r2._r.append(it)
    r3 = paragraph.add_run()
    r3._r.append(_el("w:fldChar", **{"w:fldCharType": "separate"}))
    r4 = paragraph.add_run("按 F9 或右键“更新域”生成目录")
    r5 = paragraph.add_run()
    r5._r.append(_el("w:fldChar", **{"w:fldCharType": "end"}))


def enable_update_fields(doc):
    settings = doc.settings.element
    for tag in ("w:updateFields",):
        if settings.find(qn(tag)) is None:
            settings.append(_el(tag, **{"w:val": "true"}))


def add_page_number_footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(p, " PAGE ")


def shade(cell, hexcolor):
    cell._tc.get_or_add_tcPr().append(_el("w:shd", **{"w:val": "clear",
                                                       "w:color": "auto",
                                                       "w:fill": hexcolor}))


def style_run(run, size=None, bold=None, ea=None, ascii_=None, italic=None):
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if ea or ascii_:
        run.font.name = ascii_ or FONT_ASCII
        run._element.rPr.rFonts.set(qn("w:eastAsia"), ea or FONT_EA)


# ----------------------------------------------------------------- 行内解析
def add_inline(paragraph, text, size=None, base_bold=False):
    """支持 **加粗** 与 `等宽`。"""
    tokens = []
    pos = 0
    pattern = re.compile(r"\*\*([^*]+)\*\*|`([^`]+)`")
    for m in pattern.finditer(text):
        if m.start() > pos:
            tokens.append(("plain", text[pos:m.start()]))
        if m.group(1) is not None:
            tokens.append(("bold", m.group(1)))
        else:
            tokens.append(("code", m.group(2)))
        pos = m.end()
    if pos < len(text):
        tokens.append(("plain", text[pos:]))
    if not tokens:
        tokens.append(("plain", ""))
    for kind, chunk in tokens:
        if kind == "bold":
            # 加粗片段内部可能再嵌入 `等宽`，此处按反引号再切分一次
            parts = chunk.split("`")
            for idx, seg in enumerate(parts):
                if not seg:
                    continue
                run = paragraph.add_run(seg)
                if idx % 2 == 1:
                    style_run(run, size=(size or 10.5) - 0.5, bold=True,
                              ascii_="Consolas", ea="宋体")
                else:
                    style_run(run, size=size, bold=True)
        elif kind == "code":
            run = paragraph.add_run(chunk)
            style_run(run, size=(size or 10.5) - 0.5, ascii_="Consolas", ea="宋体")
        else:
            run = paragraph.add_run(chunk.replace("`", ""))
            style_run(run, size=size, bold=True if base_bold else None)


# ----------------------------------------------------------------- 内容块
def add_caption(doc, text, above=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    style_run(r, size=9, ea=FONT_EA_HEAD, ascii_=FONT_ASCII)
    return p


def add_image(doc, alt, rel_path):
    path = rel_path
    if not os.path.isabs(path):
        path = os.path.join(BASE, rel_path)
    if not os.path.exists(path):
        p = doc.add_paragraph()
        add_inline(p, f"［缺失插图：{rel_path}］")
        return
    from PIL import Image
    with Image.open(path) as im:
        w, h = im.size
    width_cm = min(CONTENT_WIDTH_CM, 16.5)
    if h > w * 1.15:                       # 竖长图限高
        max_h = 17.0
        if width_cm * h / w > max_h:
            width_cm = max_h * w / h
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(path, width=Cm(width_cm))
    add_caption(doc, alt)


# ----------------------------------------------------------------- 标识符断行
ZWSP = "\u200b"          # 零宽空格：只影响断行，不显示宽度
MD_ESCAPE_RE = re.compile(r"\\([\\`*_{}\[\]()#+\-.!|<>~])")


def unescape_md(text):
    """还原 Markdown 反斜杠转义，避免 \< \> \| 等原样出现在正文里。"""
    return MD_ESCAPE_RE.sub(r"\1", text)


def soft_breaks(text):
    """在标识符的自然边界插入零宽空格，使换行发生在 _ . / - 或驼峰分界处，
    而不是把 fk_tblStudentProfileApplication_student 这类名字逐字符劈开。"""
    return _add_breaks(unescape_md(text))


TRAILING_PUNCT = set("、。，；：！？）》」』】…—％℃)]}>")
# 只在标识符的分段符与驼峰分界处断行；括号、尖括号、逗号等保留在标识符内部，
# 使 VARCHAR(36)、List<TermView>、DECIMAL(12,2) 这类短类型名整体不折行。
BREAK_AFTER = r"_./\\-"


def _add_breaks(s):
    """在标识符的自然边界插入零宽空格，作为 Word 的优先断行点。"""
    s = re.sub(r"(?<=[" + BREAK_AFTER + r"])(?=[0-9A-Za-z])", ZWSP, s)
    s = re.sub(r"(?<=[a-z0-9])(?=[A-Z])", ZWSP, s)
    return s


def _plain_text(text):
    """去掉行内标记，用于估算列的显示宽度。"""
    return unescape_md(text).replace("**", "").replace("`", "")


def _display_units(text):
    """估算字符串的显示宽度：全角字符按 2 个单位，半角按 1 个。"""
    n = 0
    for ch in _plain_text(text):
        if ch == ZWSP:
            continue
        n += 2 if unicodedata.east_asian_width(ch) in ("W", "F") else 1
    return n


def _longest_segment_units(text):
    """最长"不可断行片段"的宽度。

    零宽空格、空格、全角字符以及句读标点之后都是允许的断点；但句读、右括号
    等尾随标点按中文禁则不能出现在行首，因此计入前一段，并在其后断行。
    """
    plain = _plain_text(text).replace(ZWSP, "\x00")
    best = 0
    cur = 0
    for ch in plain:
        if ch.isspace() or ch == "\x00":
            best = max(best, cur)
            cur = 0
        elif ch in TRAILING_PUNCT:
            cur += 2                       # 与前面的文字同排
            best = max(best, cur)          # 标点之后可以断行
            cur = 0
        elif unicodedata.east_asian_width(ch) in ("W", "F"):
            best = max(best, cur)
            cur = 0
            best = max(best, 2)            # 单个全角字符自成一段
        else:
            cur += 1
    return max(best, cur)


CHAR_CM = 0.167          # 9.5pt 下一个半角字符的近似宽度（cm）
CELL_PAD_CM = 0.36       # 单元格左右内边距 + 少量余量


def allocate_widths(rows, ncols, total_cm, min_cm=1.6):
    """分配列宽：先给每列留出"容纳最长不可断行片段"的宽度，再把余量按内容量分配。

    第一段保证 edu.seu.vcampus.client.course.ui、
    fk_tblStudentProfileApplication_student 这类长标识符不会碎成竖排阶梯；
    第二段把剩下的宽度按内容总量分给承载说明文字的长列。
    """
    totals = []
    floors = []
    floor_cap = total_cm * 0.26
    for j in range(ncols):
        total = 0.0
        seg = 0.0
        for row in rows:
            if j < len(row):
                total += _display_units(row[j])
                seg = max(seg, _longest_segment_units(row[j]))
        totals.append(max(total, 8.0))
        need = seg * CHAR_CM + CELL_PAD_CM
        floors.append(min(max(need, min_cm), max(floor_cap, min_cm)))

    floor_sum = sum(floors)
    if floor_sum >= total_cm:                     # 极窄表：按比例压缩
        k = total_cm / floor_sum
        return [f * k for f in floors]

    remaining = total_cm - floor_sum
    total_weight = sum(totals)
    widths = [floors[j] + remaining * totals[j] / total_weight
              for j in range(ncols)]

    # 均衡修正：把宽度从"行数明显更少"的列挪给"最差单元格最高"的列，
    # 避免某列因为个别很长的单元格而被压出很高的行。
    def worst_lines(ws):
        out = []
        for j in range(ncols):
            usable = max(ws[j] - CELL_PAD_CM, 0.2)
            worst = 0
            for row in rows:
                if j < len(row):
                    units = _display_units(row[j])
                    worst = max(worst, units * CHAR_CM / usable)
            out.append(worst)
        return out

    step = 0.12
    for _ in range(240):
        lines = worst_lines(widths)
        jmax = max(range(ncols), key=lambda j: lines[j])
        donors = [j for j in range(ncols)
                  if j != jmax and lines[j] <= lines[jmax] - 1.5
                  and widths[j] - floors[j] > 0.35]
        if not donors:
            break
        jmin = max(donors, key=lambda j: widths[j] - floors[j])
        before = max(worst_lines(widths))
        widths[jmin] -= step
        widths[jmax] += step
        if max(worst_lines(widths)) >= before - 1e-9:
            widths[jmin] += step
            widths[jmax] -= step
            break
    return widths


TBLPR_ORDER = [
    "tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize",
    "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd",
    "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption",
    "tblDescription",
]


def _reorder_tblPr(tblPr):
    """按 OOXML schema 顺序重排 tblPr 子元素，避免 Word 报文档结构错误。"""
    index = {name: i for i, name in enumerate(TBLPR_ORDER)}
    children = list(tblPr)
    children.sort(key=lambda e: index.get(
        e.tag.split("}")[-1], len(TBLPR_ORDER)))
    for child in children:
        tblPr.append(child)


def set_cell_margins(table, left_twips=85, right_twips=85):
    """收紧单元格左右内边距，把省下的宽度让给正文。"""
    tblPr = table._tbl.tblPr
    mar = _el("w:tblCellMar")
    for side, value in (("top", 20), ("left", left_twips),
                        ("bottom", 20), ("right", right_twips)):
        mar.append(_el(f"w:{side}", **{"w:w": str(value), "w:type": "dxa"}))
    old = tblPr.find(qn("w:tblCellMar"))
    if old is not None:
        tblPr.remove(old)
    tblPr.append(mar)
    _reorder_tblPr(tblPr)


def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(_el("w:tblHeader", **{"w:val": "true"}))
    trPr.append(_el("w:cantSplit", **{"w:val": "true"}))


def set_word_wrap(paragraph):
    """关闭"允许西文在单词中间换行"。

    关闭后 Word 只在空格、零宽空格与全角字符处断行，配合 soft_breaks() 插入的
    零宽空格，长标识符会断在 _ . / 或驼峰分界处，而不是被逐字符劈成竖排阶梯。
    """
    ppr = paragraph._p.get_or_add_pPr()
    old = ppr.find(qn("w:wordWrap"))
    if old is not None:
        ppr.remove(old)
    el = _el("w:wordWrap", **{"w:val": "1"})
    anchor = None
    for tag in ("w:spacing", "w:ind", "w:jc", "w:rPr"):
        found = ppr.find(qn(tag))
        if found is not None:
            anchor = found
            break
    if anchor is not None:
        anchor.addprevious(el)
    else:
        ppr.append(el)


def add_table(doc, rows):
    # 先在标识符边界插入零宽空格，再据此测量与渲染
    rows = [[soft_breaks(c) for c in row] for row in rows]
    header, body = rows[0], rows[1:]
    ncols = len(header)
    t = doc.add_table(rows=1, cols=ncols)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER

    # 1) 先写入内容，再据内容宽度分配列宽
    hdr = t.rows[0].cells
    for i, txt in enumerate(header):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = 1.05
        add_inline(p, txt, size=9.5, base_bold=True)
        set_word_wrap(p)
        shade(hdr[i], "DCE9F7")
        hdr[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    for row in body:
        cells = t.add_row().cells
        for i in range(ncols):
            txt = row[i] if i < len(row) else ""
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.05
            # 纯编号列居中，其余左对齐
            if i == 0 and _display_units(txt) <= 4:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(p, txt, size=9.5)
            set_word_wrap(p)
            cells[i].vertical_alignment = WD_ALIGN_VERTICAL.TOP

    # 2) 固定表格布局并按内容比例设置列宽
    t.autofit = False
    widths = allocate_widths(rows, ncols, CONTENT_WIDTH_CM)
    for i, wcm in enumerate(widths):
        t.columns[i].width = Cm(wcm)
    for row in t.rows:
        for i, cell in enumerate(row.cells):
            cell.width = Cm(widths[i])

    set_cell_margins(t)
    set_repeat_header(t.rows[0])
    return t


def add_code_block(doc, code_lines):
    for ln in code_lines:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        shade_paragraph(p, "F4F6F8")
        r = p.add_run(ln if ln.strip() else " ")
        style_run(r, size=9, ascii_="Consolas", ea="宋体")
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)


def shade_paragraph(paragraph, hexcolor):
    ppr = paragraph._p.get_or_add_pPr()
    ppr.append(_el("w:shd", **{"w:val": "clear", "w:color": "auto",
                               "w:fill": hexcolor}))


def split_row(line):
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|") and not line.endswith("\\|"):
        line = line[:-1]
    return [c.strip() for c in re.split(r"(?<!\\)\|", line)]


# ----------------------------------------------------------------- 主转换
def render_markdown(doc, text, first_chapter=False):
    lines = text.splitlines()
    i = 0
    pending_table = []

    def flush_table():
        nonlocal pending_table
        if pending_table:
            add_table(doc, pending_table)
            pending_table = []
            doc.add_paragraph().paragraph_format.space_after = Pt(2)

    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        stripped = line.strip()

        if stripped.startswith("```"):
            i += 1
            block = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i].rstrip())
                i += 1
            i += 1
            flush_table()
            add_code_block(doc, block)
            continue

        if stripped.startswith("|") and stripped.endswith("|"):
            cells = split_row(stripped)
            if TABLE_SEP_RE.match(stripped):
                i += 1
                continue
            pending_table.append(cells)
            i += 1
            continue
        flush_table()

        if not stripped:
            i += 1
            continue

        m = IMG_RE.match(stripped)
        if m:
            add_image(doc, m.group(1), m.group(2))
            i += 1
            continue

        if stripped.startswith("#### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(3)
            add_inline(p, stripped[5:], size=10.5, base_bold=True)
            i += 1
            continue

        if stripped.startswith("### "):
            h = doc.add_heading(level=3)
            add_inline(h, stripped[4:], size=12)
            i += 1
            continue

        if stripped.startswith("## "):
            h = doc.add_heading(level=2)
            add_inline(h, stripped[3:], size=14)
            i += 1
            continue

        if stripped.startswith("# "):
            if not first_chapter:
                doc.add_page_break()
            h = doc.add_heading(level=1)
            add_inline(h, stripped[2:], size=16)
            i += 1
            continue

        if stripped.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run(stripped.lstrip("> ").strip())
            style_run(r, size=9.5, italic=True)
            i += 1
            continue

        if re.match(r"^[-*] ", stripped):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.9)
            p.paragraph_format.first_line_indent = Cm(-0.45)
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, "●　" + stripped[2:], size=10.5)
            i += 1
            continue

        if re.match(r"^\d+[.、] ", stripped):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.75)
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, stripped, size=10.5)
            i += 1
            continue

        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0.74)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.3
        add_inline(p, stripped, size=10.5)
        i += 1
    flush_table()


# ----------------------------------------------------------------- 封面
def build_cover(doc):
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("虚拟校园系统")
    style_run(r, size=30, bold=True, ea=FONT_EA_HEAD)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(30)
    r = p.add_run("软件设计说明书")
    style_run(r, size=26, bold=True, ea=FONT_EA_HEAD)

    for _ in range(5):
        doc.add_paragraph()

    for label, value in (("撰稿人：", "　"), ("版本号：", "V1.0"),
                         ("日　期：", "2026-09-11")):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(8)
        r = p.add_run(f"{label}{value}")
        style_run(r, size=14)

    for _ in range(6):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("东南大学计算机科学与工程学院")
    style_run(r, size=15, bold=True, ea=FONT_EA_HEAD)

    # 修改记录（用普通段落标题，避免进入目录）
    doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    style_run(p.add_run("文档修改记录"), size=16, bold=True, ea=FONT_EA_HEAD)
    add_table(doc, [
        ["修改日期", "版本号", "修改人", "修改说明"],
        ["2026-09-11", "V1.0", "项目组", "首次发布，覆盖用户管理、学籍管理、选课系统、图书馆、校园商城五个必做模块。"],
        ["", "", "", ""],
        ["", "", "", ""],
    ])
    doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    style_run(p.add_run("目　　录"), size=16, bold=True, ea=FONT_EA_HEAD)
    p = doc.add_paragraph()
    add_field(p, r'TOC \o "1-3" \h \z \u')


# ----------------------------------------------------------------- main
def main():
    if not os.path.isdir(CHAPTER_DIR):
        print("缺少章节目录：", CHAPTER_DIR)
        return 1
    files = sorted(glob.glob(os.path.join(CHAPTER_DIR, "*.md")))
    if not files:
        print("章节目录为空：", CHAPTER_DIR)
        return 1

    doc = Document(DEMO)
    # 清空 DEMO 正文，仅保留节属性
    body = doc.element.body
    for child in list(body):
        if child.tag == qn("w:sectPr"):
            continue
        body.remove(child)

    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(PAGE_MARGIN_CM)
    sec.top_margin = sec.bottom_margin = Cm(PAGE_MARGIN_CM)
    add_page_number_footer(sec)

    set_style_font(doc.styles["Normal"], size=10.5)
    try:
        set_style_font(doc.styles["Heading 1"], size=16, bold=True,
                       ea_font=FONT_EA_HEAD, color=(0x1F, 0x3B, 0x5C))
        set_style_font(doc.styles["Heading 2"], size=14, bold=True,
                       ea_font=FONT_EA_HEAD, color=(0x1F, 0x3B, 0x5C))
    except KeyError:
        pass
    ensure_heading3(doc)
    set_style_font(doc.styles["Heading 3"], size=12, bold=True,
                   ea_font=FONT_EA_HEAD, color=(0x2E, 0x4E, 0x6E))

    build_cover(doc)

    for idx, f in enumerate(files):
        with open(f, encoding="utf-8") as fh:
            render_markdown(doc, fh.read(), first_chapter=False)
        print("  + 已并入章节：", os.path.basename(f))

    enable_update_fields(doc)
    doc.save(OUT)
    print("已生成：", OUT)
    return 0


if __name__ == "__main__":
    sys.exit(main())
